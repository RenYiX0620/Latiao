"""Write text content to a file. Creates parent directories if needed."""
import os

_BLOCKED_DIRS = ("/etc", "/System", "/usr", "/bin", "/sbin", "/var", "/private/etc")
# 敏感目录/文件名清单统一走 cmd_safety 单点（read_file.py 早已如此）。此前这里留了
# 5 项本地拷贝，与 fallback 的 9 项漂移：~/.netrc、.git-credentials、.env.local
# 在本路径（实际生效的那条）曾放行（2026-09-24 复查发现）。
from cmd_safety import _BLOCKED_DIR_SUBSTRINGS as _BLOCKED_SUBSTRINGS  # noqa: F401
from cmd_safety import sensitive_write_name_block

NAME = "write_file"
PERMISSION = "confirm"

DEFINITION = {
    "type": "function",
    "function": {
        "name": "write_file",
        "description": "Write text to a file (parent dirs created). Path ending in .docx creates a REAL Word document, .pdf creates a REAL PDF, .xlsx creates a REAL Excel workbook — all from Markdown (headings, lists, **bold**, | tables |) — use this for reports, do NOT write a .py generator script. ⚠️ Requires user confirmation.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Absolute file path. End with .docx for Word, .xlsx for Excel."},
                "content": {"type": "string", "description": "File text. For .docx: Markdown body (title as # heading). For .xlsx: Markdown tables or TSV rows."}
            },
            "required": ["path", "content"]
        }
    }
}


def _safe_path(path: str) -> str | None:
    """返回规范化后的绝对路径；不合法返回 None。"""
    if not path:
        return None
    expanded = os.path.expanduser(path)
    if not os.path.isabs(expanded):
        # 相对路径按 sidecar 工作目录解析（模型常发 "." 或相对路径，
        # 此前直接拒绝并误报"路径穿越"）
        try:
            from agent_loop import _safe_cwd
            base = _safe_cwd()
            if not base:
                return None
            expanded = os.path.join(base, expanded)
        except Exception:
            return None
    # Block path traversal — 两种分隔符都查
    if ".." in path.split("/") or ".." in path.split("\\"):
        return None
    # realpath 解析符号链接，防止经由 symlink 逃出校验
    return os.path.realpath(expanded)




def _md_to_docx_zip(text: str) -> bytes:
    """把 Markdown 风格正文打成**合法 .docx**（纯标准库，不依赖 python-docx）。

    支持：# 标题 1–3 级、- / * 列表、空行分段、**粗体**（无样式近似为正文）。
    """
    import zipfile
    import io
    import re as _re2

    def esc(s: str) -> str:
        return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    paras = []
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            paras.append(("Heading3", esc(line[4:].strip())))
        elif line.startswith("## "):
            paras.append(("Heading2", esc(line[3:].strip())))
        elif line.startswith("# "):
            paras.append(("Heading1", esc(line[2:].strip())))
        elif line.startswith("#### "):
            paras.append(("Heading3", esc(line[5:].strip())))
        elif _re2.match(r"^\s*[-*] ", line):
            paras.append(("ListParagraph", esc(_re2.sub(r"^\s*[-*] ", "", line))))
        else:
            # 去掉 **粗体** 星号
            body = _re2.sub(r"\*\*(.+?)\*\*", r"\1", line.strip())
            paras.append(("Normal", esc(body)))
    if not paras:
        paras = [("Normal", "")]

    body_xml = []
    for style, txt in paras:
        if style == "ListParagraph":
            txt = "• " + txt
        body_xml.append(
            f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
            f'<w:r><w:t xml:space="preserve">{txt}</w:t></w:r></w:p>'
        )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{''.join(body_xml)}"
        '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr>'
        "</w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
    return buf.getvalue()


_DOCX_LATIN_FONT = "Arial"   # macOS/Windows 默认都有；中文另有 eastAsia（见 _docx_cjk_font）


def _style_profile(text: str) -> dict:
    """按**正文语言**选报告规范（不看 UI 语言：中文用户也会写英文报告）。

    中文报告：标题居中、正文小四(12pt)、**首行缩进 2 字符**、1.5 倍行距、段间距 0；
    日文：同上但缩进 1 字符；其它（英/俄…）：无缩进、段后 6pt、1.15 倍行距。
    """
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    kana = any("\u3040" <= ch <= "\u30ff" for ch in text)
    latin = sum(1 for ch in text if ch.isascii() and ch.isalpha())
    if kana:
        return {"body_pt": 12, "indent_pt": 12, "space_after": 0, "line": 1.5}
    if cjk >= 20 and cjk * 4 >= latin:
        return {"body_pt": 12, "indent_pt": 24, "space_after": 0, "line": 1.5}
    return {"body_pt": 11, "indent_pt": 0, "space_after": 6, "line": 1.15}


def _docx_cjk_font() -> str:
    """按平台挑一个该平台默认就有的中文字体名（英文名，避免区域差异）。
    口径：Word **按字体名引用、不嵌字体** → 写该平台一定存在的名字最稳；
    "跨端完全一致"由 PDF 路径（reportlab + 随包 OFL 字体）承担。
    """
    import platform
    if platform.system() == "Windows":
        return "Microsoft YaHei"
    if platform.system() == "Darwin":
        return "PingFang SC"
    return "Noto Sans SC"


def _apply_fonts(run, cjk: str) -> None:
    """中文字体必须落到 w:eastAsia —— 只设 font.name 对中文不生效（走主题字体）。"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run.font.name = _DOCX_LATIN_FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi"):
        rFonts.set(qn(attr), _DOCX_LATIN_FONT)
    rFonts.set(qn("w:eastAsia"), cjk)


def _is_numeric(cell: str) -> bool:
    """数字列（含千分位/百分号/正负号）→ 右对齐。"""
    import re as _rn
    s = cell.strip().replace(",", "").replace("%", "").replace("+", "")
    return bool(s) and bool(_rn.fullmatch(r"-?\d+(\.\d+)?", s))


def _md_to_docx(text: str) -> bytes | None:
    """Markdown → 带完整样式的 .docx（python-docx）。库缺失返回 None，由调用方降级。"""
    try:
        import io
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return None

    cjk = _docx_cjk_font()
    doc = Document()

    # 页面：A4 + 2.5cm 边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(2.5))

    # 报告规范（按正文语言）：中文=小四(12pt)/首行缩进 2 字符/1.5 倍行距/段间距 0
    prof = _style_profile(text)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(prof["body_pt"])
    normal.font.name = _DOCX_LATIN_FONT
    normal.paragraph_format.space_after = Pt(prof["space_after"])
    normal.paragraph_format.line_spacing = prof["line"]
    if prof["indent_pt"]:
        normal.paragraph_format.first_line_indent = Pt(prof["indent_pt"])
    _nrf = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi"):
        _nrf.set(qn(attr), _DOCX_LATIN_FONT)
    _nrf.set(qn("w:eastAsia"), cjk)
    for name, size, align in (("Heading 1", 22, WD_ALIGN_PARAGRAPH.CENTER),
                              ("Heading 2", 16, None),
                              ("Heading 3", 14, None)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
        st.paragraph_format.space_before = Pt(8)
        # 文档标题（# 一级）居中：中文报告习惯；二级/三级保持左对齐
        st.paragraph_format.space_after = Pt(12 if align is not None else 6)
        if align is not None:
            st.paragraph_format.alignment = align
        _srf = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi"):
            _srf.set(qn(attr), _DOCX_LATIN_FONT)
        _srf.set(qn("w:eastAsia"), cjk)

    # 页脚页码（PAGE 域）
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fld = OxmlElement("w:fldSimple")
    _fld.set(qn("w:instr"), "PAGE")
    footer_p._p.append(_fld)

    def add_par(text_line: str, style: str | None = None, indent: bool = False):
        par = doc.add_paragraph(style=style)
        if indent:
            par.paragraph_format.left_indent = Cm(0.75)
        if style is not None or indent:
            # 标题/列表/引用不吃正文的"首行缩进 2 字符"（那是正文段落的规矩）
            par.paragraph_format.first_line_indent = Pt(0)
        for i, seg in enumerate(text_line.split("**")):   # 行内 **粗体**
            if not seg:
                continue
            r = par.add_run(seg)
            if i % 2 == 1:
                r.bold = True
            _apply_fonts(r, cjk)
        return par

    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        # 表格：| a | b | + 分隔行
        if line.startswith("|") and line.endswith("|") and i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if nxt.startswith("|") and set(nxt) <= set("|-: "):
                rows: list[list[str]] = []
                j = i
                while j < len(lines) and lines[j].strip().startswith("|"):
                    cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                    if not (cells and set("".join(cells)) <= set("-: ")):
                        rows.append(cells)
                    j += 1
                if rows:
                    width = len(rows[0])
                    t = doc.add_table(rows=len(rows), cols=width)
                    t.style = "Table Grid"
                    t.autofit = False
                    numeric_cols = {
                        c for c in range(width)
                        if all(_is_numeric(r[c]) for r in rows[1:] if c < len(r)) and len(rows) > 1
                    }
                    for ri, row in enumerate(rows):
                        for ci, cell in enumerate(row[:width]):
                            tc = t.cell(ri, ci)
                            par = tc.paragraphs[0]
                            par.paragraph_format.first_line_indent = Pt(0)  # 单元格不缩进
                            r = par.add_run(cell)
                            _apply_fonts(r, cjk)
                            if ri == 0:
                                r.bold = True
                                shd = OxmlElement("w:shd")
                                shd.set(qn("w:fill"), "D9E2F3")
                                tc._tc.get_or_add_tcPr().append(shd)
                            elif ci in numeric_cols:
                                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    trPr = t.rows[0]._tr.get_or_add_trPr()   # 表头跨页重复
                    trPr.append(OxmlElement("w:tblHeader"))
                    doc.add_paragraph()
                i = j
                continue
        if line.startswith("#"):
            lvl = min(len(line) - len(line.lstrip("#")), 3)
            add_par(line.lstrip("#").strip(), style=f"Heading {lvl}")
        elif line.startswith(("- ", "* ", "+ ")):
            add_par(line[2:].strip(), style="List Bullet")
        elif len(line) > 2 and line[0].isdigit() and line[1] in ".、)":
            add_par(line[2:].strip(), style="List Number")
        elif line.startswith("> "):
            par = add_par(line[2:].strip(), indent=True)
            for r in par.runs:
                r.italic = True
        elif len(line) >= 3 and set(line) <= set("-—="):
            pass                                          # 分割线
        else:
            add_par(line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_PDF_FONT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "fonts")


def _pdf_fonts() -> tuple[str, str, bool]:
    """(regular, bold, 是否内嵌)。优先随包 OFL TTF（Noto Sans SC，可再分发）；
    没有就退回 reportlab 内置 CID 中文字体（不内嵌，靠阅读器替换，字形随阅读器）。"""
    from reportlab.pdfbase import pdfmetrics
    reg = os.path.join(_PDF_FONT_DIR, "NotoSansSC-Regular.ttf")
    bold = os.path.join(_PDF_FONT_DIR, "NotoSansSC-Bold.ttf")
    if os.path.exists(reg):
        try:
            from reportlab.pdfbase.ttfonts import TTFont
            pdfmetrics.registerFont(TTFont("LatiaoCJK", reg))
            if os.path.exists(bold):
                pdfmetrics.registerFont(TTFont("LatiaoCJK-Bold", bold))
                return "LatiaoCJK", "LatiaoCJK-Bold", True
            return "LatiaoCJK", "LatiaoCJK", True
        except Exception:
            pass
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light", "STSong-Light", False


def _md_to_pdf(text: str) -> bytes | None:
    """Markdown → PDF（reportlab）。库缺失返回 None；样式沿用 _style_profile 的语言档案。"""
    try:
        import io
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        return None

    reg_font, bold_font, embedded = _pdf_fonts()
    prof = _style_profile(text)
    body_pt = prof["body_pt"]
    leading = body_pt * prof["line"]
    hs = {"1": 22, "2": 16, "3": 14}
    st_h = {k: ParagraphStyle(f"h{k}", fontName=bold_font, fontSize=v, leading=v * 1.3,
                              alignment=TA_CENTER if k == "1" else 0,
                              spaceBefore=8, spaceAfter=12 if k == "1" else 6,
                              textColor=colors.HexColor("#1F2A37")) for k, v in hs.items()}
    st_body = ParagraphStyle("body", fontName=reg_font, fontSize=body_pt, leading=leading,
                             firstLineIndent=prof["indent_pt"], spaceAfter=prof["space_after"],
                             wordWrap="CJK")
    st_list = ParagraphStyle("li", parent=st_body, firstLineIndent=0, leftIndent=14, spaceAfter=2)
    st_cell = ParagraphStyle("cell", parent=st_body, firstLineIndent=0, spaceAfter=0,
                             leftIndent=0, fontSize=body_pt - 1, leading=(body_pt - 1) * 1.3,
                             wordWrap="CJK")
    st_cell_r = ParagraphStyle("cellr", parent=st_cell, alignment=TA_RIGHT)
    st_head = ParagraphStyle("th", parent=st_cell, fontName=bold_font)

    def esc(s: str) -> str:
        s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        parts = s.split("**")                     # 行内 **粗体** → <b>
        return "".join(f"<b>{p}</b>" if i % 2 else p for i, p in enumerate(parts))

    flow = []
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and line.endswith("|") and i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if nxt.startswith("|") and set(nxt) <= set("|-: "):
                rows: list[list[str]] = []
                j = i
                while j < len(lines) and lines[j].strip().startswith("|"):
                    cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                    if not (cells and set("".join(cells)) <= set("-: ")):
                        rows.append(cells)
                    j += 1
                if rows:
                    width = len(rows[0])
                    num_cols = {c for c in range(width)
                                if len(rows) > 1 and all(_is_numeric(r[c]) for r in rows[1:] if c < len(r))}
                    data = []
                    for ri, row in enumerate(rows):
                        cells = []
                        for ci in range(width):
                            cell = row[ci] if ci < len(row) else ""
                            style = st_head if ri == 0 else (st_cell_r if ci in num_cols else st_cell)
                            cells.append(Paragraph(esc(cell), style))
                        data.append(cells)
                    t = Table(data, repeatRows=1, hAlign="LEFT")
                    t.setStyle(TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BFBFBF")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ]))
                    flow.append(t)
                    flow.append(Spacer(1, 10))
                i = j
                continue
        if line.startswith("#"):
            lvl = str(min(len(line) - len(line.lstrip("#")), 3))
            flow.append(Paragraph(esc(line.lstrip("#").strip()), st_h[lvl]))
        elif line.startswith(("- ", "* ", "+ ")):
            flow.append(Paragraph("• " + esc(line[2:].strip()), st_list))
        elif len(line) > 2 and line[0].isdigit() and line[1] in ".、)":
            flow.append(Paragraph(esc(line), st_list))
        elif line.startswith("> "):
            flow.append(Paragraph("<i>" + esc(line[2:].strip()) + "</i>", st_list))
        elif len(line) >= 3 and set(line) <= set("-—="):
            pass
        else:
            flow.append(Paragraph(esc(line), st_body))
        i += 1

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(bold_font if embedded else reg_font, 9)
        canvas.drawCentredString(A4[0] / 2.0, 1.4 * cm, str(canvas.getPageNumber()))
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, topMargin=2.5 * cm, bottomMargin=2.5 * cm,
                      leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                      title="Latiao", author="Latiao").build(flow, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def _sheet_name(title: str, n: int) -> str:
    import re as _rs
    name = _rs.sub(r'[\\/*?:\[\]]', "", title).strip()
    return name[:31] or f"Sheet{n}"


def _tables_to_xlsx(text: str, dest: str) -> str:
    """Markdown 表格 / TSV → 真 .xlsx（openpyxl，随包自带）。

    多个表格 → 多个工作表；表格前的 # 标题用作表名。
    """
    from openpyxl import Workbook

    def split_row(line: str):
        s = line.strip()
        if not s or s.startswith("#"):
            return None
        if s.startswith("|") or s.count("|") >= 2:
            return [c.strip() for c in s.strip("|").split("|")]
        if "\t" in s:
            return [c.strip() for c in s.split("\t")]
        return None

    def is_sep(line: str) -> bool:
        s = line.strip().replace("|", "").replace(" ", "").replace(":", "")
        return bool(s) and all(c in "-+" for c in s)

    tables: list[tuple[str, list[list[str]]]] = []
    title = ""
    cur: list[list[str]] = []

    def flush():
        nonlocal cur
        if cur:
            tables.append((title, cur))
            cur = []

    for raw in text.replace("\r\n", "\n").split("\n"):
        if raw.strip().startswith("#"):
            flush()
            title = raw.strip().lstrip("#").strip()
            continue
        if is_sep(raw):
            continue
        row = split_row(raw)
        if row is None:
            flush()
            continue
        cur.append(row)
    flush()

    if not tables:
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
        tables = [("", [[ln] for ln in lines] or [[""]])]

    wb = Workbook()
    wb.remove(wb.active)
    for i, (name, rows) in enumerate(tables, 1):
        ws = wb.create_sheet(_sheet_name(name, i))
        for row in rows:
            ws.append(row)

    wb.save(dest)
    total = sum(len(r) for _, r in tables)
    return f"{len(tables)} 个工作表，共 {total} 行"


def execute(args: dict) -> str:
    content = args.get("content")
    if content is None:
        content = args.get("text", "")
    content = str(content)
    path_arg = args.get("path") or args.get("file") or ""
    p = _safe_path(path_arg)
    if p is None:
        return "⛔ Blocked: 路径无效（空路径或包含 .. 穿越片段）"
    # 系统目录一律拒绝写入
    if any(p == d or p.startswith(d + os.sep) for d in _BLOCKED_DIRS):
        return f"⛔ Blocked: 不允许写入系统目录 - {p}"
    # 敏感目录（密钥/凭证）一律拒绝
    if any(s in p for s in _BLOCKED_SUBSTRINGS):
        return f"⛔ Blocked: 不允许访问敏感目录 - {p}"
    # 敏感文件名一律拒绝（含 .env 家族，模板除外）——清单在 cmd_safety 单点
    _fn_block = sensitive_write_name_block(p)
    if _fn_block:
        return _fn_block
    # sidecar 的 plugins/ 目录一律拒绝（写入后下次启动会被 import 执行 = RCE）
    # 自动加载目录（extensions/skills）与应用配置：写进去的代码下次启动即执行（审计 P1 ⑧）
    try:
        from cmd_safety import sensitive_write_block
        _blk = sensitive_write_block(p)
        if _blk:
            return _blk
    except Exception:
        # fail-closed：封印不可用时拒绝写入（此前 pass = 放行，审计复查）
        return "⛔ Blocked: 写入封印不可用，已拒绝"
    sidecar_plugins = os.path.realpath(os.path.dirname(__file__))
    if p == sidecar_plugins or p.startswith(sidecar_plugins + os.sep):
        return f"⛔ Blocked: 不允许写入插件目录 - {p}"
    # .docx：正文按 Markdown 风格打包成**合法 Word**（纯标准库 OOXML）
    # （2026-09-24：此前模型只能写 .docx.py 生成脚本，用户拿到的不是 Word）。
    # .pdf：Markdown → PDF（reportlab；版式沿用同一套语言档案）
    if p.lower().endswith(".pdf") and not content.startswith("%PDF"):
        try:
            os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
            blob = _md_to_pdf(content)
            if blob is None:
                return ("⛔ 无法生成 PDF：reportlab 未安装（导入失败）。"
                        "可先写 .docx，再让用户用 Office 另存为 PDF。")
            with open(p, "wb") as f:
                f.write(blob)
            return (
                f"✅ 已生成 PDF：{p}（{len(blob)} 字节｜引擎：reportlab）\n"
                "支持 Markdown：# 标题（一级居中）、- 列表、**加粗**、| 表格 |。"
            )
        except Exception as e:
            return f"错误：生成 PDF 失败：{e}"
    if p.lower().endswith(".docx") and not content.startswith("PK\x03\x04"):
        try:
            os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
            # 排版引擎：python-docx 优先（真样式/真表格/中文字体），
            # 库缺失 → 现有裸 OOXML 降级；结果里注明用的是哪个引擎（可验证信号）
            blob = None
            engine = "python-docx"
            try:
                blob = _md_to_docx(content)
            except Exception:
                blob = None
            if blob is None:
                engine = "裸 OOXML（python-docx 不可用，已降级）"
                blob = _md_to_docx_zip(content)
            with open(p, "wb") as f:
                f.write(blob)
            return (
                f"✅ 已生成 Word 文档：{p}（{len(blob)} 字节，"
                f"{content.count(chr(10)) + 1} 行正文 → .docx｜引擎：{engine}）\n"
                "支持 Markdown：# 标题（1–3 级）、- 列表、1. 有序列表、**加粗**、"
                "| 表格 |（自动表头/数字右对齐/跨页重复）。"
            )
        except Exception as e:
            return f"错误：生成 .docx 失败：{e}"
    # .xlsx：Markdown 表格 / TSV → 真 Excel（openpyxl，随包自带）
    if p.lower().endswith(".xlsx") and not content.startswith("PK\x03\x04"):
        try:
            os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
            info = _tables_to_xlsx(content, p)
            return (
                f"✅ 已生成 Excel 表格：{p}（{info}）\n"
                "支持 Markdown 表格、TSV；多个表格 → 多个工作表，表格前 # 标题作表名。"
            )
        except Exception as e:
            return f"错误：生成 .xlsx 失败：{e}"
    import re as _re
    _bin_name = _re.search(r"\.(xls|pptx?|pdf|odt|ods|odp)(\.|$)", p, _re.I)
    if _bin_name and not content.startswith("PK\x03\x04"):
        return (
            f"⛔ {p} 是 {(_bin_name.group(1))} 二进制格式，write_file 不能直接写文本。\n"
            "Word（.docx）/ Excel（.xlsx）已支持直接写正文；\n"
            "演示文稿请写 gen.py 后 run_cmd 执行，交付真正文件。"
        )
    try:
        os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
        with open(p, "w", encoding="utf-8") as f:
            f.write(content)
        return f"✅ 已写入：{p}（{len(content)} 字符）"
    except Exception as e:
        return f"错误：{e}"
