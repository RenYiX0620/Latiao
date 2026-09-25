"""上传 / 语音识别 / TTS 路由（含上传 helper，避免反向依赖 api_routes）。

从 api_routes.py 拆出（2026-09-24）。APIRouter 由 api_routes.include_router 挂载。
"""
from __future__ import annotations

import asyncio
import base64
import io
import json
import logging
import os
import re

from pathlib import Path

import httpx
from fastapi import APIRouter, File, Query, Request, UploadFile
from fastapi.responses import JSONResponse, Response

import tts_service
from config import CONFIG_FILE, MAX_UPLOAD_SIZE  # 走 config，避免枢纽 import

logger = logging.getLogger("latiao-sidecar")

from http_json import _json_body  # noqa: E402 — 唯一定义

router = APIRouter()


def _translate_to_english(text: str) -> str:
    """用户偏好（config.upload_text_en=true）：上传的文字部分以英文上传。

    检测到中文时，用已配置的云端模型依次尝试翻译（如 GLM 配额耗尽 429
    自动换 deepseek）；未配置云端/全部失败/无中文时保留原文（fail-open）。
    """
    if not text or not re.search(r"[\u4e00-\u9fff]", text):
        return text
    logger.debug("upload translate: 检测到中文，尝试云端英化 (len=%d)", len(text))
    try:
        cfg = json.loads((Path.home() / ".local-ai-os" / "config.json").read_text(encoding="utf-8"))
        if not cfg.get("upload_text_en"):
            return text
        models = cfg.get("cloud_models", []) or []
        prompt = (
            "Translate the following text into English. Keep code blocks, tables and "
            "formatting unchanged. Do NOT change any numbers, dates or units. "
            "Output only the translation:\n\n" + text[:6000]
        )
        for m in models:
            try:
                url = (m.get("endpoint") or "").rstrip("/") + "/chat/completions"
                resp = httpx.post(
                    url,
                    headers={"Authorization": "Bearer " + str(m.get("key") or "")},
                    json={"model": m.get("name"),
                          "messages": [{"role": "user", "content": prompt}],
                          "max_tokens": 4096},
                    timeout=120, follow_redirects=True,
                )
                if resp.status_code != 200:
                    logger.info("upload translate: %s -> HTTP %s, try next", m.get("name"), resp.status_code)
                    continue
                out = (resp.json().get("choices") or [{}])[0].get("message", {}).get("content", "")
                if out and out.strip():
                    return out.strip()
            except Exception as e:
                logger.warning("upload translate via %s failed: %s", m.get("name"), e)
                continue
        return text
    except Exception as e:
        logger.warning("upload translate failed, keep original: %s", e)
        return text


def _extract_xlsx_text(content: bytes) -> str:
    """Excel（xlsx/xls）→ 表格文本（按 sheet 分行、tab 分隔）。

    上传的 xlsx 是 zip 二进制，直接 decode 会产生乱码（09-05 事故），
    必须解析为表格文本再交给模型。
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            out.append("### Sheet: " + ws.title)
            for row in ws.iter_rows(values_only=True):
                cells = [("" if v is None else str(v)) for v in row]
                if any(c.strip() for c in cells):
                    out.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(out)


def _process_upload_bytes(content: bytes, filename: str, content_type: str, translate: bool = True) -> dict:
    """统一处理上传字节：图片→base64；PDF→提取文字；xlsx→表格；文本→读取。
    content 统一经 _translate_to_english（用户偏好：文字部分英化）。"""
    filename = filename or ""
    is_image = content_type.startswith("image/")
    is_pdf = content_type == "application/pdf" or filename.lower().endswith(".pdf")
    is_xlsx = filename.lower().endswith((".xlsx", ".xls")) \
        or content_type in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            "application/vnd.ms-excel")

    if is_image:
        return {
            "status": "success",
            "content": f"图片已上传: {filename}",
            "filename": filename,
            "is_image": True,
            "base64_data": base64.b64encode(content).decode("utf-8"),
            "content_type": content_type,
            "size": len(content),
        }
    if is_pdf:
        reader = None
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            pdf_text = "\n\n".join(pages)
            if not pdf_text.strip():
                pdf_text = "(PDF 中没有可提取的文字，可能是扫描件或图片型 PDF)"
        except Exception as e:
            pdf_text = f"(PDF 解析失败: {e})"
        return {
            "status": "success",
            "content": _translate_to_english(pdf_text) if translate else pdf_text,
            "filename": filename,
            "is_pdf": True,
            "page_count": len(reader.pages) if reader is not None else 0,
            "size": len(content),
        }
    if is_xlsx:
        try:
            xlsx_text = _extract_xlsx_text(content)
            return {
                "status": "success",
                "content": _translate_to_english(xlsx_text) if translate else xlsx_text,
                "filename": filename,
                "is_xlsx": True,
                "size": len(content),
            }
        except Exception as e:
            return {"status": "error", "message": f"Excel 解析失败: {e}"}
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("latin-1", errors="replace")
    return {
        "status": "success",
        "content": _translate_to_english(text) if translate else text,
        "filename": filename,
        "is_image": False,
        "size": len(content),
    }



@router.post("/v1/upload_local")
async def upload_local(request: Request):
    """本地路径上传：Tauri 原生拖放事件（onDragDropEvent）拿到的是文件路径，
    由 sidecar 直接读盘（免前端读字节），解析逻辑与 /v1/upload_file 一致。"""
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"status": "error", "message": "无效请求"}, status_code=400)
    path = str(body.get("path") or "")
    translate = bool(body.get("translate", True))
    if not path or not os.path.isfile(path):
        return JSONResponse({"status": "error", "message": f"文件不存在: {path}"}, status_code=400)
    try:
        size = os.path.getsize(path)
        if size > MAX_UPLOAD_SIZE:
            return JSONResponse(
                {"status": "error", "message": f"文件过大 (上限 {MAX_UPLOAD_SIZE / 1024 / 1024:.0f}MB)"},
                status_code=413)
        with open(path, "rb") as f:
            content = f.read(MAX_UPLOAD_SIZE + 1)
        if len(content) > MAX_UPLOAD_SIZE:
            return JSONResponse({"status": "error", "message": "文件过大"}, status_code=413)
        import mimetypes
        mime, _ = mimetypes.guess_type(path)
        return _process_upload_bytes(content, os.path.basename(path), mime or "application/octet-stream", translate=translate)
    except Exception as e:
        return {"status": "error", "message": str(e)}


@router.post("/v1/upload_file")
async def upload_file(request: Request, file: UploadFile = File(...)):
    """文件上传：图片转 base64，PDF 提取文本，Excel 解析表格，文本直接读取；
    全部经 _process_upload_bytes 统一处理（含文字英化偏好）。"""
    cl = request.headers.get("content-length", "")
    if cl.isdigit() and int(cl) > MAX_UPLOAD_SIZE:
        return JSONResponse(
            {"status": "error", "message": f"文件过大 (上限 {MAX_UPLOAD_SIZE / 1024 / 1024:.0f}MB)"},
            status_code=413,
        )
    try:
        content = await file.read(MAX_UPLOAD_SIZE + 1)
        if len(content) > MAX_UPLOAD_SIZE:
            return JSONResponse(
                {"status": "error", "message": f"文件过大 (上限 {MAX_UPLOAD_SIZE / 1024 / 1024:.0f}MB)"},
                status_code=413,
            )
        return _process_upload_bytes(content, file.filename or "", file.content_type or "")
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ── Whisper model cache (lazy-load once, reuse across requests) ──
_whisper_model = None

_WHISPER_DIR = Path.home() / ".local-ai-os" / "whisper-tiny"


def _get_whisper_model():
    global _whisper_model
    if _whisper_model is None:
        from faster_whisper import WhisperModel
        # 优先本地目录加载: huggingface.co 在国内不可达(502),huggingface_hub
        # 的 httpx 也不吃 SSL_CERT_FILE 环境变量。已预下载模型文件到
        # ~/.local-ai-os/whisper-tiny/(镜像下载,见部署脚本),离线可用。
        local_model = _WHISPER_DIR / "model.bin"
        if local_model.exists():
            _whisper_model = WhisperModel(str(_WHISPER_DIR), device="cpu", compute_type="int8")
        else:
            # fallback: 尝试镜像在线下载
            os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
            _whisper_model = WhisperModel("tiny", device="cpu", compute_type="int8")
    return _whisper_model


# ── 文件预览（图/PDF/代码）：仅绝对路径 + 白名单扩展 + 体积上限 ──
_PREVIEW_BIN_EXT = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".pdf",
}
_PREVIEW_TEXT_EXT = {
    ".txt", ".md", ".json", ".js", ".ts", ".tsx", ".jsx", ".py", ".rs", ".go",
    ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".rb", ".php", ".swift", ".kt",
    ".html", ".css", ".scss", ".yml", ".yaml", ".toml", ".ini", ".sh", ".sql",
    ".vue", ".svelte", ".xml", ".csv",
}
_PREVIEW_MAX_BYTES = 8 * 1024 * 1024

def _find_soffice() -> str | None:
    """定位 LibreOffice soffice（Office→PDF 预览用）。优先 LATIAO_SOFFICE / 常见路径。"""
    import shutil
    for key in ("LATIAO_SOFFICE", "MIMO_SOFFICE", "SOFFICE"):
        cand = os.environ.get(key) or ""
        if cand and Path(cand).exists():
            return cand
    hit = shutil.which("soffice")
    if hit:
        return hit
    for c in (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/opt/homebrew/bin/soffice",
        str(Path.home() / "Library/Application Support/Xiaomi MiMo/runtimes/libreoffice"),
    ):
        pc = Path(c)
        if pc.is_file():
            return str(pc)
        if pc.is_dir():
            # runtimes 盅目录：搜一层 darwin-arm64/**/soffice
            for hit in pc.glob("**/MacOS/soffice"):
                return str(hit)
    return None




def _resolve_office_deliverable(p: Path) -> Path | None:
    """脚本旁的真文档：`生成报告.py` → `报告.docx` / 同目录最像的 office 文件。

    模型常写「生成脚本 + 跑出 docx」，前端点审阅拿的是脚本路径。
    服务端必须能解析到真正交付物，否则预览永远是代码。
    """
    office_exts = (".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt", ".pdf", ".odt", ".ods", ".odp")
    stem = p.name
    lower = stem.lower()
    # 1) 去掉 .py / .docx.py 等后缀，直接拼 office 扩展名
    bases = [p.with_suffix("")]
    if lower.endswith(".py"):
        bases.append(p.with_suffix("").with_suffix(""))  # report.docx.py → report.docx 再去后缀
    # 2) 去掉「生成/gen/generate」前缀
    import re as _re
    stripped = _re.sub(r"^(生成|生成器|gen[_-]?|generate[_-]?)", "", p.stem, flags=_re.I)
    if stripped and stripped != p.stem:
        bases.append(p.parent / stripped)
    for base in bases:
        for ext in office_exts:
            cand = base.with_name(base.name + ext) if base.suffix else base.parent / (base.name + ext)
            # with_suffix 在无后缀时行为：Path("报告").with_suffix(".docx") → 报告.docx
            try:
                cand = Path(str(base) + ext) if not base.name.lower().endswith(ext) else base
                if cand.is_file() and cand.stat().st_size > 100:
                    head = cand.read_bytes()[:4]
                    if head.startswith(b"PK\x03") or head.startswith(b"%PDF") or head.startswith(b"\xd0\xcf"):
                        return cand
            except OSError:
                continue
    # 3) 同目录里名字含共同子串的 office 文件（mtime 最新优先）
    try:
        key = _re.sub(r"[^\w\u4e00-\u9fff]", "", stripped or p.stem)[:6]
        hits = []
        for f in p.parent.iterdir():
            if not f.is_file() or f.suffix.lower() not in office_exts:
                continue
            if f.name == p.name:
                continue
            try:
                if f.stat().st_size < 100:
                    continue
                head = f.read_bytes()[:4]
                if not (head.startswith(b"PK\x03") or head.startswith(b"%PDF") or head.startswith(b"\xd0\xcf")):
                    continue
            except OSError:
                continue
            norm = _re.sub(r"[^\w\u4e00-\u9fff]", "", f.stem)
            score = 2 if key and key in norm else 1
            hits.append((score, f.stat().st_mtime, f))
        if hits:
            hits.sort(key=lambda x: (x[0], x[1]), reverse=True)
            return hits[0][2]
    except OSError:
        pass
    return None


_OFFICE_EXT = {".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".odt", ".ods", ".odp", ".rtf"}
_PDF_CACHE = Path.home() / ".cache" / "latiao" / "office-preview"


def _office_to_pdf(src: Path) -> Path:
    """LibreOffice 无头转 PDF。缓存：同名+mtime 命中则不重转。"""
    import subprocess
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("missing_soffice")
    src = src.resolve()
    _PDF_CACHE.mkdir(parents=True, exist_ok=True)
    stamp = f"{int(src.stat().st_mtime)}"
    out = _PDF_CACHE / f"{src.stem}.{stamp}.pdf"
    if out.exists() and out.stat().st_size > 0:
        return out
    # 独立 UserInstallation：并行/残留 soffice 进程会抢 profile 锁
    profile = _PDF_CACHE / f"profile-{stamp}-{os.getpid()}"
    try:
        from cmd_safety import child_env
        proc = subprocess.run(
            [
                soffice, "--headless", "--norestore", "--nolockcheck",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "pdf", "--outdir", str(_PDF_CACHE),
                str(src),
            ],
            capture_output=True, text=True, timeout=60,
            env=child_env(),  # 不把 LATIAO_AUTH_TOKEN / 云密钥传给 soffice
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        raise RuntimeError(f"convert_failed:{e}") from e
    finally:
        # 输出名可能是原名.pdf，再规范化到 stamp 名
        try:
            produced = _PDF_CACHE / f"{src.stem}.pdf"
            if produced.exists() and produced != out:
                if out.exists():
                    out.unlink()
                produced.rename(out)
        except OSError:
            pass
        # 清 profile 目录（小体积，可留；这里直接忽略失败）
    if not out.exists() or out.stat().st_size == 0:
        tail = (proc.stderr or proc.stdout or "")[-300:]
        raise RuntimeError(f"convert_failed:{tail}")
    return out


_PREVIEW_MAX_TEXT = 200_000


@router.get("/v1/file/preview")
async def file_preview(path: str = Query(..., description="绝对路径")):
    """预览已落盘文件：图/PDF 回 base64，代码/文本回截断正文。

    安全：仅绝对路径、禁止 ..、扩展白名单、拒绝敏感目录、体积上限。
    """
    from fastapi import Query as _Q  # noqa: F401 — Query 已在顶部导入
    raw = (path or "").strip()
    if not raw or not raw.startswith(("/", "X:", "x:")) or ".." in raw:
        return {"status": "error", "message": "invalid path"}
    p = Path(raw).expanduser().resolve()
    try:
        if not p.is_absolute() or not p.is_file():
            return {"status": "error", "message": "not a file"}
        low = str(p).lower()
        blocked_names = ("config.json", "id_rsa", "id_ed25519", "credentials",
                         "shadow", "sudoers", "secrets")
        if any(n in low for n in (".ssh", ".aws", ".gnupg", "/etc/shadow", "/etc/sudoers")):
            return {"status": "error", "message": "blocked"}
        if any(n in p.name.lower() for n in blocked_names):
            return {"status": "error", "message": "blocked"}
    except OSError:
        return {"status": "error", "message": "invalid path"}

    ext = p.suffix.lower()
    size = p.stat().st_size
    if size > _PREVIEW_MAX_BYTES:
        return {"status": "error", "message": "file too large"}
    # 点到生成器脚本时：解析同目录真文档，预览以交付物为准
    if ext in (".py", ".js", ".sh", ".ts", ".txt", ".md") or ext == "":
        doc = _resolve_office_deliverable(p)
        if doc is not None:
            p = doc
            ext = p.suffix.lower()
            size = p.stat().st_size
    # 按文件头嗅探：扩展名可能是 .docx.py / 无扩展名，但内容才是真相
    head = b""
    try:
        with open(p, "rb") as _bf:
            head = _bf.read(8)
    except OSError:
        head = b""
    if head.startswith(b"%PDF"):
        ext = ".pdf"
    elif head.startswith(b"PK\x03\x04"):
        # zip 族：docx/xlsx/pptx/odt… 一律按 office 转 PDF
        ext = ".docx"
    if ext in _PREVIEW_BIN_EXT:
        import base64
        data = p.read_bytes()
        mime = {
            ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
            ".gif": "image/gif", ".webp": "image/webp", ".svg": "image/svg+xml",
            ".pdf": "application/pdf",
        }[ext]
        return {
            "status": "ok", "kind": "pdf" if ext == ".pdf" else "image",
            "mime": mime, "name": p.name,
            "data_base64": base64.b64encode(data).decode("ascii"),
        }
    if ext in _PREVIEW_TEXT_EXT or ext == "":
        try:
            text = p.read_text(encoding="utf-8", errors="replace")[:_PREVIEW_MAX_TEXT]
        except OSError as e:
            return {"status": "error", "message": str(e)}
        return {"status": "ok", "kind": "code", "name": p.name, "text": text,
                "truncated": size > _PREVIEW_MAX_TEXT}
    # Office：LibreOffice 无头转 PDF 后走 PDF 预览；缺依赖则告知前端「系统打开」
    if ext in _OFFICE_EXT:
        import base64
        try:
            pdf_path = _office_to_pdf(p)
            data = pdf_path.read_bytes()
            return {
                "status": "ok", "kind": "pdf", "mime": "application/pdf",
                "name": p.name, "converted_from": "office",
                "data_base64": base64.b64encode(data).decode("ascii"),
            }
        except RuntimeError as e:
            msg = str(e)
            if "missing_soffice" in msg:
                return {"status": "ok", "kind": "office", "name": p.name,
                        "need_soffice": True}
            return {"status": "error", "message": f"office 转 PDF 失败：{msg[:160]}"}
    return {"status": "ok", "kind": "office", "name": p.name, "need_soffice": True}


@router.post("/v1/recognize_speech")
async def recognize_speech(request: Request):
    """语音识别：前端发 WAV base64 → faster-whisper 本地识别"""
    import tempfile

    try:
        body = await _json_body(request)
        audio_base64 = body.get("audio_base64", "")

        if not audio_base64:
            return {"status": "error", "message": "No audio data provided"}

        # 解码前先按 base64 长度做大小检查（约 25MB 原始音频上限）
        if len(audio_base64) > 25 * 1024 * 1024 * 4 // 3:
            return {"status": "error", "message": "音频过大（上限约 25MB）"}

        audio_bytes = base64.b64decode(audio_base64)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".webm") as tmp:
            tmp.write(audio_bytes)
            wav_path = tmp.name

        try:
            # 模型加载与推理都是 CPU 密集阻塞操作，放到线程执行避免卡住事件循环
            model = await asyncio.to_thread(_get_whisper_model)
            segments, info = await asyncio.to_thread(model.transcribe, wav_path, language="zh", beam_size=5)

            text = " ".join(s.text.strip() for s in segments)

            if text:
                return {"status": "success", "text": text}
            else:
                return {"status": "success", "text": "(未识别到语音内容)"}

        finally:
            if os.path.exists(wav_path):
                os.unlink(wav_path)
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ── 语音合成（朗读）：代理给独立的本地语音服务，服务不在就返回结构化降级 ──
# 第一轮「框架先行」：模型不在 Latiao 里，换模型只改 config.json 的 tts.model_id。


@router.post("/v1/synthesize_speech")
async def synthesize_speech(request: Request):
    """文字 → 音频。成功回音频字节；服务未就绪回 {code: tts_unavailable} 让前端回退系统语音。"""
    body = await _json_body(request)
    text = str(body.get("text") or "").strip()
    if not text:
        return JSONResponse(status_code=400, content={"status": "error",
                                                     "code": "tts_empty_text",
                                                     "message": "没有需要朗读的文本"})
    audio, ctype, err = await tts_service.synthesize(
        text, body.get("voice"), body.get("speed"), CONFIG_FILE,
        pitch=body.get("pitch"), emotion=body.get("emotion"))
    if audio:
        return Response(content=audio, media_type=ctype or "audio/wav")
    return JSONResponse(status_code=503, content=err or {"status": "error",
                                                        "code": "tts_unavailable"})


@router.get("/v1/tts/voices")
async def tts_voices():
    """音色列表：按当前模型动态取，前端不用因换模型而改。"""
    data = await tts_service.list_voices(CONFIG_FILE)
    code = 503 if data.get("status") == "error" else 200
    return JSONResponse(status_code=code, content=data)


@router.get("/v1/tts/status")
async def tts_status():
    """给设置页看的一行状态（是否启用 / 服务是否在跑 / 当前模型与音色）。"""
    return tts_service.status(CONFIG_FILE)


