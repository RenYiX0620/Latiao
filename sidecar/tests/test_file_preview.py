"""交付物预览接口：路径白名单与体积/类型（2026-09-24）。"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("LATIAO_AUTH_TOKEN", "preview-token")


@pytest.fixture()
def client(tmp_path, monkeypatch):
    import main as m
    m.AUTH_TOKEN = "preview-token"
    from fastapi.testclient import TestClient
    return TestClient(m.app, raise_server_exceptions=False)


@pytest.fixture()
def headers():
    return {"X-Latiao-Token": "preview-token"}


def test_preview_code_file(client, headers, tmp_path):
    f = tmp_path / "hello.py"
    f.write_text("print('hi')\n", encoding="utf-8")
    r = client.get("/v1/file/preview", params={"path": str(f)}, headers=headers)
    data = r.json()
    assert data["status"] == "ok" and data["kind"] == "code"
    assert "print" in data["text"]


def test_preview_rejects_relative_and_dotdot(client, headers):
    for bad in ("etc/passwd", "../secret", "C:foo", ""):
        r = client.get("/v1/file/preview", params={"path": bad}, headers=headers)
        assert r.json()["status"] == "error"


def test_preview_blocks_secrets(client, headers, tmp_path):
    f = tmp_path / "config.json"
    f.write_text('{"key":"sk-xxx"}', encoding="utf-8")
    r = client.get("/v1/file/preview", params={"path": str(f)}, headers=headers)
    assert r.json()["status"] == "error"

    ssh = tmp_path / ".ssh"
    ssh.mkdir()
    k = ssh / "id_rsa"
    k.write_text("PRIVATE", encoding="utf-8")
    r = client.get("/v1/file/preview", params={"path": str(k)}, headers=headers)
    assert r.json()["status"] == "error"


def test_preview_png_base64(client, headers, tmp_path):
    import base64
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    f = tmp_path / "x.png"
    f.write_bytes(png)
    r = client.get("/v1/file/preview", params={"path": str(f)}, headers=headers)
    data = r.json()
    assert data["status"] == "ok" and data["kind"] == "image"
    assert data["data_base64"]


def test_preview_office_kind(client, headers, tmp_path, monkeypatch):
    import api_routes_media as m
    f = tmp_path / "a.docx"
    f.write_bytes(b"PK\x03\x04fake")
    # 缺 LibreOffice 时应标 need_soffice，而不是去转真实文件
    monkeypatch.setattr(m, "_find_soffice", lambda: None)
    r = client.get("/v1/file/preview", params={"path": str(f)}, headers=headers)
    data = r.json()
    assert data["status"] == "ok" and data["kind"] == "office" and data.get("need_soffice")


def test_find_soffice_env_override(tmp_path, monkeypatch):
    import api_routes_media as m
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\n")
    fake.chmod(0o755)
    monkeypatch.setenv("LATIAO_SOFFICE", str(fake))
    assert m._find_soffice() == str(fake)


def test_office_to_pdf_cached(tmp_path, monkeypatch):
    import api_routes_media as m
    src = tmp_path / "a.docx"
    src.write_bytes(b"PK\x03\x04fake")
    cache = tmp_path / "cache"
    monkeypatch.setattr(m, "_PDF_CACHE", cache)
    called = {"n": 0}

    def fake_run(cmd, **kw):
        called["n"] += 1
        # 模拟 soffice 输出 a.pdf
        cache.mkdir(parents=True, exist_ok=True)
        (cache / "a.pdf").write_bytes(b"%PDF-1.4 fake")
        class R:
            returncode = 0
            stderr = ""
            stdout = "convert"
        return R()

    monkeypatch.setattr(m, "_find_soffice", lambda: "/fake/soffice")
    monkeypatch.setattr("subprocess.run", fake_run)
    out1 = m._office_to_pdf(src)
    assert out1.exists() and out1.read_bytes().startswith(b"%PDF")
    out2 = m._office_to_pdf(src)
    assert out2 == out1
    assert called["n"] == 1, "mtime 缓存命中时不应重转"


def test_office_to_pdf_missing_soffice(tmp_path, monkeypatch):
    import api_routes_media as m
    import pytest
    src = tmp_path / "a.docx"
    src.write_bytes(b"x")
    monkeypatch.setattr(m, "_find_soffice", lambda: None)
    with pytest.raises(RuntimeError, match="missing_soffice"):
        m._office_to_pdf(src)


def test_preview_office_needs_soffice_flag(client, headers, tmp_path, monkeypatch):
    import api_routes_media as m
    f = tmp_path / "b.docx"
    f.write_bytes(b"PK\x03\x04")
    monkeypatch.setattr(m, "_find_soffice", lambda: None)
    r = client.get("/v1/file/preview", params={"path": str(f)}, headers=headers)
    data = r.json()
    assert data["status"] == "ok" and data["kind"] == "office" and data.get("need_soffice")


def test_write_file_docx_text_becomes_real_word():
    """.docx + 文本 = 直接生成合法 Word（不再是 .docx.py 脚本）。"""
    import zipfile
    import plugins.write_file as wf
    dest = Path("/tmp/latiao-ut-report.docx")
    out = wf.execute({"path": str(dest), "content": "# 标题\n正文"})
    assert "✅" in out and dest.exists()
    assert zipfile.is_zipfile(dest)


def test_write_file_xlsx_markdown_tables_become_real_excel(tmp_path):
    """.xlsx + Markdown 表格 = 真 Excel；多表 → 多工作表，# 标题作表名。"""
    from openpyxl import load_workbook
    import plugins.write_file as wf
    dest = tmp_path / "行情.xlsx"
    md = (
        "# A股\n"
        "| 代码 | 名称 |\n|------|------|\n| 600519 | 贵州茅台 |\n\n"
        "# 指数\n"
        "| 名称 | 点位 |\n|------|------|\n| 上证 | 3200 |\n"
    )
    out = wf.execute({"path": str(dest), "content": md})
    assert "✅" in out and dest.exists()
    wb = load_workbook(dest)
    assert wb.sheetnames == ["A股", "指数"]
    assert wb["A股"]["A1"].value == "代码"
    assert wb["A股"]["B2"].value == "贵州茅台"
    assert wb["指数"]["B2"].value == "3200"


def test_write_file_xlsx_tsv(tmp_path):
    from openpyxl import load_workbook
    import plugins.write_file as wf
    dest = tmp_path / "tsv.xlsx"
    out = wf.execute({"path": str(dest), "content": "a\tb\n1\t2"})
    assert "✅" in out and dest.exists()
    ws = load_workbook(dest).active
    assert ws["A1"].value == "a" and ws["B2"].value == "2"


def test_write_file_xls_still_rejected(tmp_path):
    import plugins.write_file as wf
    out = wf.execute({"path": str(tmp_path / "old.xls"), "content": "x"})
    assert "⛔" in out


def test_preview_sniffs_zip_as_office(tmp_path, monkeypatch):
    import api_routes_media as m
    f = tmp_path / "weird.docx.py"
    # 伪 zip 头
    f.write_bytes(b"PK\x03\x04" + b"rest")
    monkeypatch.setattr(m, "_find_soffice", lambda: None)
    # 不按 .py 走代码；应按 office
    # 但扩展是 .py，嗅探后 ext 变 .docx 再进 office 分支
    from fastapi.testclient import TestClient
    import main as mi
    mi.AUTH_TOKEN = "preview-token"
    c = TestClient(mi.app, raise_server_exceptions=False)
    r = c.get("/v1/file/preview", params={"path": str(f)},
              headers={"X-Latiao-Token": "preview-token"})
    data = r.json()
    assert data["status"] == "ok"
    assert data["kind"] in ("office", "pdf")


def test_resolve_deliverable_strips_gen_prefix(tmp_path):
    """`生成A股分析报告.py` 旁有 `A股分析报告.docx` 时必须解析到 docx。"""
    import zipfile
    import api_routes_media as m
    import io
    doc = tmp_path / "A股分析报告_2026-09-24.docx"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w/>")
    doc.write_bytes(buf.getvalue())
    script = tmp_path / "生成A股分析报告_2026-09-24.py"
    script.write_text("print(1)", encoding="utf-8")
    hit = m._resolve_office_deliverable(script)
    assert hit is not None and hit.name == "A股分析报告_2026-09-24.docx"
    assert zipfile.is_zipfile(doc)
